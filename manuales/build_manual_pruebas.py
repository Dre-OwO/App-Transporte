from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_manual_tecnico import (
    BLUE,
    INK,
    MUTED,
    add_bullet,
    add_code,
    add_heading,
    add_numbered,
    add_para,
    add_table,
    apply_run_font,
    set_paragraph_spacing,
    setup_document,
)


OUT = Path(__file__).with_name("Manual_Pruebas_App_Transporte.docx")


def add_metadata_table(doc):
    rows = [
        ("Tipo de proyecto", "Sistema web académico/administrativo"),
        ("Nombre del proyecto", "App Transporte"),
        ("Documento", "Manual de Pruebas"),
        ("Fase", "Fase 2"),
        ("Versión", "1.0"),
        ("Fecha", "20 de mayo de 2026"),
        ("Firma de quien valida", "Pendiente de firma"),
    ]
    add_table(doc, ["Campo", "Detalle"], rows, [2600, 6760], font_size=10)


def add_team_table(doc):
    rows = [
        ("Jefe de equipo", "miembro 1"),
        ("Integrante", "miembro 1"),
        ("Integrante", "miembro 2"),
        ("Integrante", "miembro 3"),
        ("Integrante", "miembro 4"),
        ("Integrante", "miembro 5"),
    ]
    add_table(doc, ["Rol", "Nombre"], rows, [2600, 6760], font_size=10)


def add_cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(72)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("FASE 2")
    apply_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Manual de Pruebas")
    apply_run_font(run, size=28, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("App Transporte - Sistema de suscripción anual para transporte público")
    apply_run_font(run, size=13, color=MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    add_metadata_table(doc)

    add_heading(doc, "Equipo de trabajo", 2)
    add_team_table(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=0)
    r = p.add_run(
        "Documento de planeación, ejecución y evidencia de pruebas funcionales para la aplicación."
    )
    apply_run_font(r, size=10.5, color=MUTED, italic=True)


def add_contents(doc):
    add_heading(doc, "Contenido", 1, page_break_before=True)
    sections = [
        "1. Introducción",
        "2. Proyecto",
        "3. Objetivos del manual de pruebas",
        "4. Entorno y configuración de pruebas",
        "5. Base de datos a utilizar",
        "6. Criterios de aprobación/rechazo",
        "7. Datos de prueba",
        "8. Matriz de pruebas",
        "9. Ejecución de pruebas",
        "10. Registro de evidencias",
        "11. Conclusión",
    ]
    for item in sections:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=3)
        r = p.add_run(item)
        apply_run_font(r, size=10.5)


def add_test_case(doc, case_id, module, name, description, precondition, data, steps, expected):
    add_heading(doc, f"Caso de prueba {case_id}", 3)
    rows = [
        ("Módulo", module),
        ("Nombre", name),
        ("Descripción", description),
        ("Precondición", precondition),
        ("Datos de prueba", data),
        ("Pasos", "\n".join(f"{idx}. {step}" for idx, step in enumerate(steps, start=1))),
        ("Resultado esperado", expected),
        ("Resultado obtenido", "Pendiente de ejecución. Adjuntar captura o descripción del resultado obtenido."),
        ("Estado", "Pendiente / Aprobado / Rechazado"),
    ]
    add_table(doc, ["Campo", "Detalle"], rows, [2100, 7260], font_size=9.5)


def build():
    doc = Document()
    setup_document(doc)
    add_cover(doc)
    add_contents(doc)

    add_heading(doc, "1. Introducción", 1, page_break_before=True)
    add_para(
        doc,
        "El presente manual define la estrategia de pruebas para App Transporte, una aplicación web desarrollada "
        "con Flask y MongoDB para gestionar usuarios, planes de suscripción anual, suscripciones y validaciones "
        "de viaje. El documento adapta el formato base proporcionado y lo orienta a los módulos reales del proyecto.",
    )

    add_heading(doc, "2. Proyecto", 1)
    add_table(
        doc,
        ["Elemento", "Detalle"],
        [
            ("Tipo de proyecto", "Sistema web académico/administrativo"),
            ("Nombre del proyecto", "App Transporte"),
            ("Objetivo del sistema", "Administrar suscripciones anuales de transporte público y validar viajes asociados a usuarios registrados."),
            ("Tecnologías", "Python 3, Flask 3.1.0, PyMongo 4.11.3, MongoDB, HTML, CSS y Jinja2."),
            ("Módulos principales", "Inicio, registro, login, planes, panel de usuario, validación de viaje, administración y reportes."),
        ],
        [2500, 6860],
    )

    add_heading(doc, "3. Objetivos del manual de pruebas", 1)
    for item in (
        "Definir el entorno mínimo necesario para ejecutar pruebas funcionales del sistema.",
        "Establecer criterios claros para clasificar errores graves, medios y leves.",
        "Documentar casos de prueba para módulos públicos, autenticación, validación de viajes y administración.",
        "Registrar resultados esperados y dejar espacio para evidencias de ejecución.",
        "Facilitar la revisión de calidad antes de la entrega académica del proyecto.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "4. Entorno y configuración de pruebas", 1)
    add_para(
        doc,
        "Para ejecutar las pruebas se requiere un ambiente local con la aplicación Flask y MongoDB activos. "
        "Las pruebas pueden realizarse desde el mismo equipo que actúa como servidor o desde un navegador en la misma máquina.",
    )
    add_table(
        doc,
        ["Elemento", "Configuración requerida"],
        [
            ("Sistema operativo del servidor", "Windows 10+, macOS 12+ o Linux actual."),
            ("Características del equipo servidor", "Procesador Intel Core i3/Ryzen 3 o superior, 4 GB RAM mínimo, 1 GB libre en disco."),
            ("Equipos del cliente", "Equipo con navegador moderno: Chrome, Edge, Firefox o Safari."),
            ("Características de equipos de prueba", "Resolución suficiente para visualizar formularios y tablas; conexión local al servidor Flask."),
            ("Servidor de aplicación", "Flask ejecutándose en http://127.0.0.1:5000/."),
            ("Dependencias", "Entorno virtual activo e instalación de requirements.txt."),
        ],
        [3000, 6360],
    )
    add_para(doc, "Comandos sugeridos para preparar el ambiente:")
    add_code(
        doc,
        [
            "source app_transporte/bin/activate",
            "pip install -r requirements.txt",
            'export MONGO_URI=\"mongodb://localhost:27017/\"',
            'export MONGO_DB_NAME=\"app_transporte\"',
            "python3 scripts/create_admin.py",
            "python3 app.py",
        ],
    )

    add_heading(doc, "5. Base de datos a utilizar", 1)
    add_table(
        doc,
        ["Campo", "Valor"],
        [
            ("Base de datos", "app_transporte"),
            ("Servidor", "localhost"),
            ("Puerto", "27017"),
            ("Usuario", "Sin usuario en ambiente local; se recomienda app_transporte_app si se habilita autenticación."),
            ("Contraseña", "Sin contraseña en ambiente local; definir una contraseña segura si se habilita autenticación."),
            ("Colecciones", "usuarios, planes, suscripciones y validaciones."),
        ],
        [2500, 6860],
    )

    add_heading(doc, "6. Criterios de aprobación/rechazo", 1)
    add_table(
        doc,
        ["Severidad", "Descripción", "Criterio de rechazo"],
        [
            ("Errores graves", "Caída de la aplicación, corrupción de datos, autenticación incorrecta, permisos administrativos expuestos o incumplimiento de funciones principales.", "La entrega se rechaza si existe al menos un error grave abierto."),
            ("Errores medios", "Errores en presentación de datos importantes, validaciones incompletas, redirecciones incorrectas o fallas en funciones secundarias.", "La entrega se rechaza si existen errores medios sin documentar o sin plan de corrección."),
            ("Errores leves", "Detalles visuales, textos inconsistentes, pequeñas dificultades de operación o comportamientos correctos pero poco claros.", "La entrega puede aprobarse si los errores leves no bloquean el flujo principal y quedan registrados."),
        ],
        [1800, 4800, 2760],
        font_size=9.5,
    )
    add_para(
        doc,
        "Criterio general de aprobación: todos los casos críticos de autenticación, permisos, suscripción y validación "
        "de viaje deben aprobarse. Los casos no ejecutados deben quedar justificados en el registro de evidencias.",
        bold_prefix="Criterio general de aprobación:",
    )

    add_heading(doc, "7. Datos de prueba", 1)
    add_table(
        doc,
        ["Dato", "Valor sugerido", "Uso"],
        [
            ("Administrador", "admin@app-transporte.test / admin123", "Acceso al panel administrativo y creación de datos base."),
            ("Cliente válido", "cliente@app-transporte.test / cliente123", "Registro, login, suscripción y validación de viaje."),
            ("Cliente sin suscripción", "sinplan@app-transporte.test / cliente123", "Validar rechazo de viaje sin suscripción activa."),
            ("Plan activo", "Plan Anual General, precio 2500, límite 4 viajes diarios", "Suscripción desde módulo público y administración."),
            ("Plan inactivo", "Plan Prueba Inactivo", "Validar que no se pueda asignar desde flujo normal."),
            ("Ruta de prueba", "Ruta 622, Terminal Norte", "Registro de validaciones de viaje."),
        ],
        [2200, 3300, 3860],
        font_size=9.5,
    )

    add_heading(doc, "8. Matriz de pruebas", 1)
    matrix_rows = [
        ("CDP-LOGIN-01", "Inicio de sesión", "Acceso con credenciales válidas", "Alta"),
        ("CDP-LOGIN-02", "Inicio de sesión", "Acceso con contraseña incorrecta", "Alta"),
        ("CDP-LOGIN-03", "Inicio de sesión", "Acceso con correo no registrado", "Alta"),
        ("CDP-LOGIN-04", "Seguridad", "Acceso a módulo protegido sin sesión", "Alta"),
        ("CDP-LOGIN-05", "Seguridad", "Cliente intenta acceder al panel administrativo", "Alta"),
        ("CDP-REG-01", "Registro", "Alta de usuario cliente válida", "Alta"),
        ("CDP-REG-02", "Registro", "Correo inválido", "Media"),
        ("CDP-REG-03", "Registro", "Contraseña menor a 6 caracteres", "Media"),
        ("CDP-REC-01", "Recuperación", "Cambio de contraseña con datos correctos", "Media"),
        ("CDP-PLAN-01", "Planes", "Consulta pública de planes activos", "Media"),
        ("CDP-SUS-01", "Suscripciones", "Cliente se suscribe a plan activo", "Alta"),
        ("CDP-SUS-02", "Suscripciones", "Cliente intenta duplicar suscripción activa", "Alta"),
        ("CDP-VAL-01", "Validación", "Validación aceptada con suscripción activa", "Alta"),
        ("CDP-VAL-02", "Validación", "Validación rechazada sin suscripción activa", "Alta"),
        ("CDP-ADM-01", "Administración", "Creación de usuario desde admin", "Media"),
        ("CDP-ADM-02", "Administración", "Creación de plan desde admin", "Media"),
        ("CDP-ADM-03", "Administración", "Suscripción con fecha final anterior", "Alta"),
        ("CDP-REP-01", "Reportes", "Carga de dashboard y reportes", "Media"),
    ]
    add_table(doc, ["ID", "Módulo", "Caso", "Prioridad"], matrix_rows, [1650, 2200, 4210, 1300], font_size=9.2)

    add_heading(doc, "9. Ejecución de pruebas", 1)
    add_test_case(
        doc,
        "CDP-LOGIN-01",
        "Inicio de sesión",
        "Acceso con credenciales válidas",
        "Verificar que el sistema permita el acceso a un usuario registrado con correo y contraseña correctos.",
        "El usuario debe existir, estar activo y tener contraseña configurada.",
        "Correo: admin@app-transporte.test. Contraseña: admin123.",
        ["Abrir http://127.0.0.1:5000/login.", "Ingresar el correo del administrador.", "Ingresar la contraseña correcta.", "Presionar el botón Entrar."],
        "El sistema debe mostrar el mensaje de bienvenida y redirigir al panel administrativo si el rol es admin.",
    )
    add_test_case(
        doc,
        "CDP-LOGIN-02",
        "Inicio de sesión",
        "Acceso con contraseña incorrecta",
        "Verificar que el sistema rechace credenciales cuando la contraseña no corresponde al usuario.",
        "El correo debe existir en la base de datos.",
        "Correo: admin@app-transporte.test. Contraseña: incorrecta123.",
        ["Abrir la pantalla de login.", "Ingresar el correo registrado.", "Ingresar una contraseña incorrecta.", "Presionar Entrar."],
        "El sistema debe denegar el acceso y mostrar el mensaje: Correo o contrasena incorrectos.",
    )
    add_test_case(
        doc,
        "CDP-LOGIN-03",
        "Inicio de sesión",
        "Acceso con correo no registrado",
        "Verificar que el sistema no permita inicio de sesión con un correo inexistente.",
        "No aplica.",
        "Correo: noexiste@app-transporte.test. Contraseña: cliente123.",
        ["Abrir la pantalla de login.", "Ingresar un correo no registrado.", "Ingresar cualquier contraseña.", "Presionar Entrar."],
        "El sistema debe denegar el acceso y mostrar el mensaje genérico: Correo o contrasena incorrectos.",
    )
    add_test_case(
        doc,
        "CDP-LOGIN-04",
        "Seguridad de sesión",
        "Acceso a módulo protegido sin sesión activa",
        "Verificar que el sistema restrinja el acceso directo a módulos internos si el usuario no ha iniciado sesión.",
        "No debe existir sesión activa en el navegador.",
        "URL de prueba: http://127.0.0.1:5000/panel-usuario.",
        ["Cerrar sesión o abrir una ventana privada.", "Ingresar directamente la URL /panel-usuario.", "Presionar Enter."],
        "El sistema debe redirigir a /login y mostrar el mensaje: Inicia sesion para continuar.",
    )
    add_test_case(
        doc,
        "CDP-LOGIN-05",
        "Seguridad de roles",
        "Cliente intenta acceder al panel administrativo",
        "Verificar que un usuario con rol cliente no pueda entrar a rutas administrativas.",
        "Debe existir sesión activa de un usuario con rol cliente.",
        "Usuario: cliente@app-transporte.test.",
        ["Iniciar sesión como cliente.", "Ingresar directamente http://127.0.0.1:5000/admin.", "Observar la redirección y mensaje."],
        "El sistema debe impedir el acceso, redirigir al panel del usuario y mostrar: No tienes permisos para entrar al panel administrativo.",
    )
    add_test_case(
        doc,
        "CDP-REG-01",
        "Registro",
        "Alta de usuario cliente válida",
        "Verificar que un visitante pueda crear una cuenta de cliente con datos válidos.",
        "El correo no debe existir previamente en la colección usuarios.",
        "Nombre: Cliente Prueba. Correo: cliente@app-transporte.test. Teléfono: 3312345678. Contraseña: cliente123.",
        ["Abrir /registro.", "Capturar nombre, correo, teléfono y contraseña.", "Presionar Crear usuario."],
        "El sistema debe crear el usuario, iniciar sesión automáticamente y mostrar: Cuenta creada correctamente. Bienvenido.",
    )
    add_test_case(
        doc,
        "CDP-REG-02",
        "Registro",
        "Correo inválido",
        "Verificar validación de formato de correo durante el registro.",
        "No aplica.",
        "Correo: correo-invalido. Contraseña: cliente123.",
        ["Abrir /registro.", "Capturar datos con correo inválido.", "Presionar Crear usuario."],
        "El sistema debe rechazar el registro y mostrar: Ingresa un correo valido.",
    )
    add_test_case(
        doc,
        "CDP-REG-03",
        "Registro",
        "Contraseña menor a 6 caracteres",
        "Verificar que el sistema no permita contraseñas demasiado cortas.",
        "El correo de prueba no debe existir previamente.",
        "Contraseña: 123.",
        ["Abrir /registro.", "Capturar datos válidos y contraseña de 3 caracteres.", "Presionar Crear usuario."],
        "El sistema debe rechazar el registro y mostrar: La contrasena debe tener al menos 6 caracteres.",
    )
    add_test_case(
        doc,
        "CDP-REC-01",
        "Recuperación de contraseña",
        "Cambio de contraseña con datos correctos",
        "Verificar que el usuario pueda actualizar su contraseña si proporciona correo, teléfono y confirmación válidos.",
        "El usuario debe existir y tener teléfono registrado.",
        "Correo: cliente@app-transporte.test. Teléfono: 3312345678. Nueva contraseña: cliente456.",
        ["Abrir /recuperar-contrasena.", "Ingresar correo, teléfono, nueva contraseña y confirmación.", "Presionar el botón del formulario."],
        "El sistema debe actualizar la contraseña, redirigir a login y mostrar: Contrasena actualizada. Ya puedes iniciar sesion.",
    )
    add_test_case(
        doc,
        "CDP-PLAN-01",
        "Planes",
        "Consulta pública de planes activos",
        "Verificar que la pantalla de planes muestre únicamente planes activos disponibles para contratación.",
        "Debe existir al menos un plan activo en la colección planes.",
        "Plan activo: Plan Anual General.",
        ["Abrir /planes sin iniciar sesión.", "Revisar la tabla de planes.", "Confirmar que se visualiza nombre, precio, límite diario y acción."],
        "El sistema debe listar planes activos. Si el usuario no tiene sesión, debe mostrar enlace para iniciar sesión antes de suscribirse.",
    )
    add_test_case(
        doc,
        "CDP-SUS-01",
        "Suscripciones",
        "Cliente se suscribe a plan activo",
        "Verificar que un cliente autenticado pueda contratar un plan activo.",
        "Debe existir sesión de cliente y al menos un plan activo. El cliente no debe tener suscripción activa.",
        "Usuario: cliente@app-transporte.test. Plan: Plan Anual General.",
        ["Iniciar sesión como cliente.", "Abrir /planes.", "Presionar Suscribirse en el plan activo."],
        "El sistema debe crear la suscripción, redirigir al panel de usuario y mostrar: Suscripcion creada correctamente.",
    )
    add_test_case(
        doc,
        "CDP-SUS-02",
        "Suscripciones",
        "Cliente intenta duplicar suscripción activa",
        "Verificar que el sistema impida registrar una segunda suscripción activa para el mismo usuario.",
        "El cliente ya debe tener una suscripción activa.",
        "Usuario: cliente@app-transporte.test. Plan: Plan Anual General.",
        ["Iniciar sesión como cliente con suscripción activa.", "Abrir /planes.", "Intentar suscribirse nuevamente a un plan activo."],
        "El sistema debe rechazar la operación y mostrar: Ya tienes una suscripcion activa.",
    )
    add_test_case(
        doc,
        "CDP-VAL-01",
        "Validación de viaje",
        "Validación aceptada con suscripción activa",
        "Verificar que el sistema registre una validación aceptada cuando el usuario tiene suscripción activa vigente.",
        "Debe existir sesión de cliente con suscripción activa y fecha_fin mayor o igual a la fecha actual.",
        "Ruta: Ruta 622. Estación: Terminal Norte.",
        ["Iniciar sesión como cliente con suscripción activa.", "Abrir /validar-viaje.", "Capturar ruta y estación.", "Presionar Registrar validacion."],
        "El sistema debe registrar la validación, mostrar estado aceptada y motivo: Suscripcion activa encontrada.",
    )
    add_test_case(
        doc,
        "CDP-VAL-02",
        "Validación de viaje",
        "Validación rechazada sin suscripción activa",
        "Verificar que el sistema registre una validación rechazada cuando el usuario no tiene suscripción activa.",
        "Debe existir sesión de cliente sin suscripción activa.",
        "Ruta: Ruta 622. Estación: Terminal Norte.",
        ["Iniciar sesión como cliente sin suscripción activa.", "Abrir /validar-viaje.", "Capturar ruta y estación.", "Presionar Registrar validacion."],
        "El sistema debe registrar la validación con resultado rechazada y explicar que el usuario no tiene una suscripción activa.",
    )
    add_test_case(
        doc,
        "CDP-ADM-01",
        "Administración de usuarios",
        "Creación de usuario desde admin",
        "Verificar que el administrador pueda crear usuarios desde el módulo administrativo.",
        "Debe existir sesión activa de administrador.",
        "Usuario nuevo: prueba.admin@app-transporte.test / prueba123 / rol cliente.",
        ["Iniciar sesión como administrador.", "Abrir /admin/usuarios.", "Capturar nombre, correo, teléfono, contraseña, rol y estado.", "Presionar Crear."],
        "El sistema debe insertar el usuario y mostrar: Usuario creado correctamente.",
    )
    add_test_case(
        doc,
        "CDP-ADM-02",
        "Administración de planes",
        "Creación de plan desde admin",
        "Verificar que el administrador pueda crear planes con precio y límite numéricos.",
        "Debe existir sesión activa de administrador.",
        "Plan: Plan QA. Precio anual: 2500. Límite diario: 4.",
        ["Iniciar sesión como administrador.", "Abrir /admin/planes.", "Capturar nombre, descripción, precio, límite y estado activo.", "Presionar Crear."],
        "El sistema debe crear el plan y mostrar: Plan creado correctamente.",
    )
    add_test_case(
        doc,
        "CDP-ADM-03",
        "Administración de suscripciones",
        "Suscripción con fecha final anterior",
        "Verificar que el administrador no pueda crear una suscripción con fecha de fin menor a fecha de inicio.",
        "Debe existir sesión activa de administrador, usuario y plan seleccionables.",
        "fecha_inicio: 2026-05-20. fecha_fin: 2026-05-19.",
        ["Iniciar sesión como administrador.", "Abrir /admin/suscripciones.", "Seleccionar usuario y plan.", "Capturar fechas inválidas.", "Presionar Crear."],
        "El sistema debe rechazar la operación y mostrar: La fecha de fin no puede ser anterior a la fecha de inicio.",
    )
    add_test_case(
        doc,
        "CDP-REP-01",
        "Reportes",
        "Carga de dashboard y reportes",
        "Verificar que el administrador pueda consultar estadísticas y reportes básicos sin errores de base de datos.",
        "Debe existir sesión activa de administrador y MongoDB encendido.",
        "Rutas: /admin y /admin/reportes.",
        ["Iniciar sesión como administrador.", "Abrir /admin.", "Abrir /admin/reportes.", "Revisar estadísticas, top de planes y validaciones recientes."],
        "El sistema debe mostrar dashboard y reportes. Si no hay datos, debe presentar valores en cero o estados vacíos sin caer.",
    )

    add_heading(doc, "10. Registro de evidencias", 1)
    add_para(
        doc,
        "Durante la ejecución se recomienda adjuntar capturas por cada caso de prueba, especialmente cuando el "
        "resultado obtenido sea diferente al esperado. La evidencia mínima debe incluir fecha, caso de prueba, "
        "usuario usado, pantalla evaluada y estado final.",
    )
    add_table(
        doc,
        ["Caso", "Fecha", "Responsable", "Evidencia", "Estado"],
        [
            ("CDP-LOGIN-01", "", "miembro 1", "", ""),
            ("CDP-LOGIN-02", "", "miembro 2", "", ""),
            ("CDP-REG-01", "", "miembro 3", "", ""),
            ("CDP-SUS-01", "", "miembro 4", "", ""),
            ("CDP-VAL-01", "", "miembro 5", "", ""),
        ],
        [1700, 1400, 1900, 3060, 1300],
        font_size=9.2,
    )

    add_heading(doc, "11. Conclusión", 1)
    add_para(
        doc,
        "El manual de pruebas proporciona una base organizada para validar que App Transporte cumpla sus flujos "
        "principales: autenticación, gestión de usuarios, consulta y contratación de planes, validación de viajes "
        "y administración. Una vez ejecutados los casos, los resultados obtenidos deberán actualizarse con capturas "
        "y observaciones para cerrar formalmente la fase de pruebas.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
