from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Manual_Tecnico_App_Transporte.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 89, 89)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "B7C7D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                cell.width = Inches(widths[idx] / 1440)
                set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_field(paragraph, field):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fld_end)


def apply_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 24, INK, 0, 8),
        ("Subtitle", 13, MUTED, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        if name.startswith("Heading"):
            style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", 1)
    code.font.name = "Courier New"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    code.font.size = Pt(9.5)
    code.font.color.rgb = RGBColor(40, 40, 40)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.1

    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("Manual Técnico | App Transporte")
    apply_run_font(run, size=9, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("Página ")
    apply_run_font(run, size=9, color=MUTED)
    add_field(footer_p, "PAGE")


def add_heading(doc, text, level, page_break_before=False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.page_break_before = page_break_before
    for run in p.runs:
        apply_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        apply_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        apply_run_font(r2)
    else:
        r = p.add_run(text)
        apply_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    apply_run_font(r)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    apply_run_font(r)
    return p


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph(style="Code Block")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        r = p.add_run(line)
        apply_run_font(r, size=9.5, name="Courier New")
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_GRAY)
        p_pr.append(shd)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    set_table_borders(table)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.1)
        r = p.add_run(header)
        apply_run_font(r, size=font_size, color=INK, bold=True)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.15)
            if idx == 0 and len(headers) <= 3:
                r = p.add_run(str(value))
                apply_run_font(r, size=font_size, color=INK, bold=True)
            else:
                r = p.add_run(str(value))
                apply_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_metadata_table(doc):
    rows = [
        ("Institución", "Universidad de Guadalajara - CUCEI"),
        ("Asignatura", "Calidad de Software"),
        ("Proyecto", "App Transporte"),
        ("Documento", "Manual Técnico"),
        ("Fase", "Fase 1"),
        ("Versión", "1.0"),
        ("Fecha", "20 de mayo de 2026"),
    ]
    add_table(doc, ["Campo", "Detalle"], rows, [2100, 7260], font_size=10)


def add_cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(80)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("FASE 1")
    apply_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Manual Técnico")
    apply_run_font(run, size=28, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("App Transporte - Sistema de suscripción anual para transporte público")
    apply_run_font(run, size=13, color=MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    add_metadata_table(doc)
    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=0)
    r = p.add_run(
        "Documento técnico de instalación, configuración, arquitectura, usuarios y contingencias del sistema."
    )
    apply_run_font(r, size=10.5, color=MUTED, italic=True)


def add_contents(doc):
    add_heading(doc, "Contenido", 1, page_break_before=True)
    sections = [
        "1. Objetivos",
        "1.1 Objetivos específicos",
        "2. Alcance",
        "3. Requerimientos técnicos",
        "3.1 Requerimientos mínimos de hardware",
        "3.2 Requerimientos mínimos de software",
        "4. Herramientas utilizadas para el desarrollo",
        "5. Instalación",
        "6. Configuración",
        "7. Diseño de la arquitectura física",
        "8. Usuarios",
        "8.1 Usuarios de base de datos",
        "8.2 Usuarios de sistemas operativos",
        "8.3 Usuarios de aplicaciones",
        "9. Contingencias y soluciones",
    ]
    for item in sections:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=3)
        r = p.add_run(item)
        apply_run_font(r, size=10.5)


def build():
    doc = Document()
    setup_document(doc)
    add_cover(doc)
    add_contents(doc)

    add_heading(doc, "1. Objetivos", 1, page_break_before=True)
    add_para(
        doc,
        "El objetivo de este manual es documentar la información técnica necesaria para instalar, configurar, "
        "ejecutar y mantener en ambiente local el proyecto App Transporte. El sistema permite gestionar usuarios, "
        "planes de suscripción anual, suscripciones activas y validaciones de viaje mediante una aplicación web "
        "desarrollada con Flask y MongoDB.",
    )

    add_heading(doc, "1.1 Objetivos específicos", 2)
    for item in (
        "Establecer los requerimientos mínimos de hardware y software para ejecutar la aplicación.",
        "Describir el proceso de instalación, configuración inicial y arranque del sistema.",
        "Identificar la arquitectura física y lógica que conecta navegador, aplicación Flask y base de datos MongoDB.",
        "Documentar los usuarios técnicos y funcionales del sistema, junto con sus privilegios principales.",
        "Registrar contingencias frecuentes de instalación, conexión y operación, así como sus soluciones recomendadas.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "2. Alcance", 1)
    add_para(
        doc,
        "Este documento está dirigido al equipo de desarrollo, al responsable de calidad de software, al profesor "
        "de la asignatura y a cualquier evaluador técnico que necesite levantar o revisar el proyecto en un equipo local.",
    )
    add_para(
        doc,
        "El manual cubre la ejecución de la versión académica del sistema, incluyendo módulos públicos, módulo "
        "administrativo, conexión a MongoDB, creación del primer administrador y flujo básico de validación de viajes.",
    )
    add_para(
        doc,
        "No cubre despliegue productivo en nube, integración con validadores físicos, pasarelas de pago reales, "
        "monitoreo avanzado ni endurecimiento completo de seguridad para producción.",
    )
    add_para(
        doc,
        "Conocimientos básicos requeridos: uso de terminal, Python, entornos virtuales, Flask, MongoDB, HTML/CSS "
        "y navegación web.",
        bold_prefix="Conocimientos básicos requeridos:",
    )

    add_heading(doc, "3. Requerimientos técnicos", 1)
    add_para(
        doc,
        "App Transporte es una aplicación web local. La carga esperada para esta etapa es baja, por lo que los "
        "requerimientos se enfocan en disponer de Python, MongoDB y un navegador moderno correctamente instalados.",
    )

    add_heading(doc, "3.1 Requerimientos mínimos de hardware", 2)
    add_table(
        doc,
        ["Recurso", "Mínimo", "Recomendado"],
        [
            ("Procesador", "Intel Core i3, AMD Ryzen 3 o equivalente", "Intel Core i5, AMD Ryzen 5 o superior"),
            ("Memoria RAM", "4 GB", "8 GB o más para ejecutar navegador, Flask y MongoDB con comodidad"),
            ("Almacenamiento", "1 GB libre para código, entorno virtual y dependencias", "5 GB libres para base de datos local y respaldos"),
            ("Red", "Interfaz local; no requiere Internet para operar una vez instaladas las dependencias", "Conexión a Internet para instalación y control de versiones"),
        ],
        [1850, 3600, 3910],
    )

    add_heading(doc, "3.2 Requerimientos mínimos de software", 2)
    add_table(
        doc,
        ["Software", "Versión/valor", "Uso en el proyecto"],
        [
            ("Sistema operativo", "Windows 10+, macOS 12+ o Linux actual", "Ejecución del entorno local de desarrollo"),
            ("Python", "3.10 o superior", "Lenguaje principal del backend"),
            ("pip / venv", "Incluidos con Python", "Instalación aislada de dependencias"),
            ("Flask", "3.1.0", "Framework web, rutas, sesiones, formularios y plantillas"),
            ("PyMongo", "4.11.3", "Conexión y operaciones CRUD contra MongoDB"),
            ("MongoDB", "Local; recomendado MongoDB Community 6.0+ u 8.0 en macOS/Homebrew", "Persistencia de usuarios, planes, suscripciones y validaciones"),
            ("Navegador web", "Chrome, Edge, Firefox o Safari reciente", "Acceso a la interfaz pública y administrativa"),
            ("Privilegios", "Usuario estándar para ejecutar; administrador si se instalará Python o MongoDB", "Preparación del ambiente"),
        ],
        [2050, 2500, 4810],
    )

    add_heading(doc, "4. Herramientas utilizadas para el desarrollo", 1)
    add_table(
        doc,
        ["Herramienta", "Descripción", "Motivo de uso"],
        [
            ("Python 3", "Lenguaje de programación principal.", "Permite implementar la lógica de negocio y conectar la aplicación con MongoDB."),
            ("Flask 3.1.0", "Framework web ligero para Python.", "Maneja rutas, sesiones, mensajes flash, formularios y renderizado de plantillas."),
            ("Jinja2", "Motor de plantillas incluido con Flask.", "Permite generar vistas HTML dinámicas para módulos públicos y administrativos."),
            ("MongoDB", "Base de datos NoSQL orientada a documentos.", "Almacena usuarios, planes, suscripciones y validaciones con estructura flexible."),
            ("PyMongo 4.11.3", "Driver oficial de MongoDB para Python.", "Ejecuta consultas, inserciones, actualizaciones, índices y agregaciones desde el backend."),
            ("HTML y CSS", "Tecnologías base de interfaz web.", "Construyen la experiencia visual de la aplicación pública y el panel administrativo."),
            ("Git", "Sistema de control de versiones.", "Registra cambios del proyecto y facilita el seguimiento de entregas."),
        ],
        [1700, 3300, 4360],
    )

    add_heading(doc, "5. Instalación", 1)
    add_para(doc, "Los pasos siguientes asumen que el código fuente se encuentra en la carpeta del proyecto.")

    steps = [
        ("Entrar a la carpeta del proyecto.", ['cd "Calidad de software"']),
        ("Crear el entorno virtual si aún no existe.", ["python3 -m venv app_transporte"]),
        ("Activar el entorno virtual.", ["source app_transporte/bin/activate", r"app_transporte\\Scripts\\activate  # Windows"]),
        ("Instalar las dependencias declaradas.", ["pip install -r requirements.txt"]),
        ("Encender MongoDB local.", ["brew services start mongodb-community@8.0  # macOS con Homebrew", 'mongosh "mongodb://localhost:27017/"  # verificación opcional']),
        ("Configurar variables de entorno recomendadas.", ['export MONGO_URI="mongodb://localhost:27017/"', 'export MONGO_DB_NAME="app_transporte"', 'export SECRET_KEY="cambiar-esta-llave-en-produccion"']),
        ("Crear el primer administrador.", ["python3 scripts/create_admin.py"]),
        ("Ejecutar la aplicación Flask.", ["python3 app.py"]),
        ("Abrir el sistema en el navegador.", ["http://127.0.0.1:5000/"]),
    ]
    for idx, (text, code_lines) in enumerate(steps, start=1):
        add_numbered(doc, f"{idx}. {text}")
        add_code(doc, code_lines)

    add_para(
        doc,
        "Recomendación: iniciar MongoDB antes de ejecutar Flask. Si MongoDB está apagado, la interfaz puede abrir, "
        "pero los módulos que consultan datos mostrarán advertencias de conexión.",
        bold_prefix="Recomendación:",
    )

    add_heading(doc, "6. Configuración", 1)
    add_para(
        doc,
        "La configuración principal se realiza mediante variables de entorno. El código también incluye valores "
        "por defecto para facilitar la ejecución local durante el desarrollo.",
    )
    add_table(
        doc,
        ["Elemento", "Valor actual o sugerido", "Descripción"],
        [
            ("SECRET_KEY", "Variable opcional; por defecto dev-secret-key", "Llave de Flask para sesiones y mensajes flash. Debe cambiarse fuera de desarrollo."),
            ("SESSION_COOKIE_HTTPONLY", "True", "Reduce exposición de la cookie de sesión a scripts del navegador."),
            ("SESSION_COOKIE_SAMESITE", "Lax", "Disminuye riesgos básicos de envío de cookies entre sitios."),
            ("MONGO_URI", "mongodb://localhost:27017/", "Cadena de conexión al servidor MongoDB."),
            ("MONGO_DB_NAME", "app_transporte", "Nombre de la base de datos utilizada por la aplicación."),
            ("Modo debug", "app.run(debug=True)", "Útil en desarrollo; debe desactivarse si se publica el sistema."),
        ],
        [2300, 2750, 4310],
    )

    add_heading(doc, "Configuración de base de datos", 2)
    add_para(doc, "Al iniciar la conexión, el módulo database.py prepara índices para mejorar búsquedas y evitar duplicados.")
    add_table(
        doc,
        ["Colección", "Índice principal", "Uso"],
        [
            ("usuarios", "correo único", "Evita cuentas duplicadas y permite autenticación por correo."),
            ("planes", "nombre único", "Evita registrar dos planes con el mismo nombre."),
            ("suscripciones", "usuario_id + estado", "Facilita localizar suscripciones activas por usuario."),
            ("validaciones", "usuario_id + fecha_hora", "Permite consultar historial reciente de validaciones."),
        ],
        [1900, 3000, 4460],
    )

    add_heading(doc, "Módulos y rutas principales", 2)
    add_table(
        doc,
        ["Módulo", "Rutas", "Función"],
        [
            ("Público", "/, /planes, /registro, /login, /recuperar-contrasena", "Consulta de planes, creación de cuenta, autenticación y recuperación de contraseña."),
            ("Usuario autenticado", "/panel-usuario, /validar-viaje, /suscribirse/<plan_id>", "Consulta de suscripción, validación de viajes y alta de suscripción."),
            ("Administrativo", "/admin, /admin/usuarios, /admin/planes, /admin/suscripciones, /admin/reportes", "Gestión de datos, dashboard y reportes básicos."),
            ("Soporte", "/logout, /reportes, /easter-egg", "Cierre de sesión, redirección a reportes y ruta adicional del proyecto."),
        ],
        [1900, 3450, 4010],
    )

    add_heading(doc, "7. Diseño de la arquitectura física", 1)
    add_para(
        doc,
        "En la fase actual, la aplicación se ejecuta como una arquitectura local de desarrollo: navegador, servidor "
        "Flask y MongoDB residen en el mismo equipo. La separación lógica de responsabilidades sí está definida.",
    )
    add_table(
        doc,
        ["Capa", "Componente", "Archivo/servicio", "Responsabilidad"],
        [
            ("Cliente", "Navegador web", "Chrome, Firefox, Edge o Safari", "Presenta formularios, tablas y navegación del sistema."),
            ("Aplicación", "Flask", "app.py", "Define rutas, sesiones, permisos, formularios y renderizado de vistas."),
            ("Datos", "Capa de acceso a datos", "database.py", "Valida entradas, ejecuta operaciones CRUD y serializa documentos."),
            ("Persistencia", "MongoDB", "Base app_transporte", "Guarda usuarios, planes, suscripciones y validaciones."),
            ("Presentación", "Plantillas y estilos", "templates/ y static/css/styles.css", "Organiza vistas públicas, vistas admin y diseño visual."),
        ],
        [1300, 1900, 2500, 3660],
        font_size=9.5,
    )
    add_table(
        doc,
        ["Elemento físico", "Detalle"],
        [
            ("Nombre de equipos", "Equipo de desarrollo local y servicio MongoDB local."),
            ("Ubicación física", "Equipo personal o laboratorio académico. No se define sitio alterno en esta fase."),
            ("Direcciones IP", "127.0.0.1 / localhost para Flask y MongoDB local."),
            ("Puertos TCP/UDP", "TCP 5000 para Flask; TCP 27017 para MongoDB."),
            ("Dependencias externas", "Python 3, Flask, PyMongo, MongoDB y navegador web."),
            ("Recomendación futura", "Separar servidor de aplicación, servidor de base de datos y respaldo periódico si se lleva a producción."),
        ],
        [2300, 7060],
    )

    add_heading(doc, "Estructura de carpetas relevante", 2)
    add_table(
        doc,
        ["Ruta", "Descripción"],
        [
            ("app.py", "Punto de entrada Flask; rutas públicas, rutas admin, sesiones y control de permisos."),
            ("database.py", "Conexión MongoDB, validaciones, CRUD, agregaciones y serialización de documentos."),
            ("requirements.txt", "Dependencias Python del proyecto: Flask y PyMongo."),
            ("scripts/create_admin.py", "Script interactivo para crear el primer administrador activo."),
            ("templates/public/", "Vistas de inicio, planes, registro, login, panel de usuario y validación de viaje."),
            ("templates/admin/", "Vistas de dashboard, usuarios, planes, suscripciones y reportes."),
            ("static/css/styles.css", "Hoja de estilos general para interfaz pública y administrativa."),
        ],
        [2550, 6810],
    )

    add_heading(doc, "8. Usuarios", 1)
    add_para(
        doc,
        "El sistema contempla usuarios técnicos para ejecutar la solución y usuarios funcionales dentro de la aplicación. "
        "La versión local simplifica la autenticación de MongoDB, pero se documenta el usuario recomendado para un escenario formal.",
    )

    add_heading(doc, "8.1 Usuarios de base de datos", 2)
    add_table(
        doc,
        ["Usuario", "Propósito", "Privilegios"],
        [
            ("Sin autenticación local", "Configuración actual esperada para desarrollo con MongoDB en localhost.", "Acceso local al servidor MongoDB según configuración del equipo."),
            ("app_transporte_app", "Usuario técnico recomendado si se habilita autenticación en MongoDB.", "readWrite sobre la base app_transporte; acceso a usuarios, planes, suscripciones y validaciones."),
        ],
        [2200, 3450, 3710],
    )

    add_heading(doc, "8.2 Usuarios de sistemas operativos", 2)
    add_table(
        doc,
        ["Usuario", "Propósito", "Privilegios sobre carpetas"],
        [
            ("Usuario local del desarrollador", "Instalar dependencias, activar entorno virtual, ejecutar MongoDB y levantar Flask.", "Lectura y escritura sobre la carpeta del proyecto y sobre app_transporte/ si el entorno virtual se crea localmente."),
            ("Administrador del equipo", "Instalar Python, MongoDB o herramientas del sistema cuando no existan.", "Privilegios temporales de instalación; no se requiere para usar la aplicación ya instalada."),
        ],
        [2350, 3600, 3410],
    )

    add_heading(doc, "8.3 Usuarios de aplicaciones", 2)
    add_table(
        doc,
        ["Rol", "Acceso", "Privilegios dentro de la aplicación"],
        [
            ("Administrador", "Panel administrativo protegido por sesión y rol admin.", "Gestiona usuarios, planes y suscripciones; consulta dashboard y reportes."),
            ("Cliente", "Interfaz pública y panel personal después de iniciar sesión.", "Se registra, consulta planes, crea suscripción, revisa su panel y registra validaciones de viaje."),
            ("Visitante", "Interfaz pública sin sesión.", "Consulta inicio y planes; puede iniciar sesión o crear una cuenta."),
        ],
        [1800, 3200, 4360],
    )

    add_heading(doc, "9. Contingencias y soluciones", 1)
    add_para(
        doc,
        "La siguiente matriz resume problemas frecuentes durante instalación, configuración u operación local del sistema.",
    )
    add_table(
        doc,
        ["Contingencia", "Causa probable", "Solución recomendada"],
        [
            ("La aplicación abre, pero no carga datos.", "MongoDB no está encendido o MONGO_URI apunta a una dirección incorrecta.", "Iniciar MongoDB y verificar conexión con mongosh. Confirmar MONGO_URI=mongodb://localhost:27017/."),
            ("Error por Flask o PyMongo faltante.", "Dependencias no instaladas o entorno virtual no activado.", "Activar app_transporte e instalar con pip install -r requirements.txt."),
            ("No se puede crear administrador inicial.", "MongoDB no responde o ya existen administradores activos.", "Encender MongoDB. Si ya existe un admin, confirmar en el script si se desea crear otro."),
            ("Correo de usuario duplicado.", "La colección usuarios tiene índice único sobre correo.", "Usar un correo diferente o editar/eliminar el usuario existente desde el panel administrativo."),
            ("Nombre de plan duplicado.", "La colección planes tiene índice único sobre nombre.", "Cambiar el nombre del plan o actualizar el registro existente."),
            ("No se puede eliminar un plan.", "El plan tiene suscripciones asociadas.", "Cancelar o eliminar las suscripciones vinculadas antes de eliminar el plan."),
            ("El usuario no puede acceder a /admin.", "La sesión no existe o el rol no es admin.", "Iniciar sesión con un usuario administrador creado mediante scripts/create_admin.py."),
            ("La validación de viaje es rechazada.", "No existe usuario, no tiene suscripción activa o la suscripción venció.", "Verificar correo, estado de la suscripción y fecha_fin en el módulo de suscripciones."),
            ("Puerto 5000 ocupado.", "Otro proceso está usando el puerto de Flask.", "Cerrar el proceso en conflicto o ejecutar Flask en un puerto alterno durante pruebas."),
            ("Se expone el modo debug o la llave dev-secret-key.", "Configuración de desarrollo usada fuera del ambiente local.", "Desactivar debug y definir SECRET_KEY segura antes de publicar el sistema."),
        ],
        [2600, 3150, 3610],
        font_size=9.2,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
